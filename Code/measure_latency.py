#!/usr/bin/env python3
"""
measure_latency.py — measures the latency numbers for the paper's latency table (Reviewer 3).
Does NOT modify agentic_app5.py.  Put it in the same folder as agentic_app5.py and run:

    python measure_latency.py                         # PatchTST inference + controller time (+ LLM time from your Excel)
    python measure_latency.py --pcap D:\\some_capture.pcap   # additionally: feature-extraction time

Output:  latency_values.tex   (copy it next to the paper's .tex file and recompile; the table fills itself)
         latency_results.json (all raw numbers)
"""
import os
import sys
import json
import time
import argparse
import statistics
import importlib.util

HERE = os.path.dirname(os.path.abspath(__file__))


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


def measure_inference(pred, csv_path, n, reps):
    """PatchTST inference: ms per window (batch of BATCH_SIZE), windows/s, and single-window latency."""
    import numpy as np
    import pandas as pd
    import torch
    import joblib
    dev = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    model = pred.PatchTSTClassifier().to(dev)
    model.load_state_dict(torch.load(pred.MODEL_PATH, map_location=dev))
    model.eval()
    scaler = joblib.load(pred.SCALER_PATH)
    df = pd.read_csv(csv_path, nrows=n)
    cols = list(scaler.feature_names_in_)
    X = scaler.transform(df[cols]).astype(np.float32).reshape(-1, pred.WINDOW, pred.BASE_FEATURES)
    Xt = torch.from_numpy(X).to(dev)
    bs = getattr(pred, "BATCH_SIZE", 256)

    def sync():
        if dev.type == "cuda":
            torch.cuda.synchronize()
    with torch.no_grad():
        for _ in range(5):                                   # warm-up
            model(Xt[:bs])
            sync()
        runs = []
        for _ in range(reps):
            t = time.perf_counter()
            for i in range(0, len(Xt), bs):
                model(Xt[i:i + bs])
            sync()
            runs.append(time.perf_counter() - t)
        single = []
        for i in range(min(200, len(Xt))):
            t = time.perf_counter()
            model(Xt[i:i + 1])
            sync()
            single.append((time.perf_counter() - t) * 1000.0)
    thr = len(Xt) / statistics.median(runs)
    return {"device": str(dev), "n_windows": len(Xt), "batch_size": bs, "windows_per_s": thr,
            "ms_per_window_batched": 1000.0 / thr, "ms_single_window_median": statistics.median(single)}


def measure_features(pred, pcap):
    """Feature extraction with predict.py's own functions (packets -> per-second features -> windows)."""
    t0 = time.perf_counter()
    raw = pred.extract_packets(pcap)
    t1 = time.perf_counter()
    feats = pred.compute_features(raw)
    t2 = time.perf_counter()
    pred.sliding_windows(feats)
    t3 = time.perf_counter()
    secs = max(1, len(feats))
    total = t3 - t0
    return {"pcap": pcap, "seconds_of_traffic": secs, "read_s": t1 - t0, "features_s": t2 - t1, "windows_s": t3 - t2,
            "total_s": total, "s_per_20s_recording": total * 20.0 / secs}


def measure_controller(app, out_dir, n_sessions=200):
    """Deterministic controller compute time per 20 s recording on THIS machine (uses the caches of your run)."""
    cdir = os.path.join(out_dir, "cache_csv")
    if not os.path.isdir(cdir):
        return None
    caches = app.load_caches(cdir)[:n_sessions]
    if not caches:
        return None
    out = {"n_sessions": len(caches)}
    for mode in ("evidence", "fast"):
        cfg = app.dc_replace(app.PolicyConfig(max_observation=app.MAX_OBSERVATION), evidence_mode=mode)
        ts = []
        for c in caches:
            t = time.perf_counter()
            app.run_policy_offline(c, app.ABLATIONS["F_full"], cfg)
            ts.append((time.perf_counter() - t) * 1000.0)
        out[f"{mode}_ms"] = statistics.fmean(ts)
    return out


def read_llm_seconds(out_dir):
    """Mean LLM wall-clock per recording from your own --auto-llm run (sheet T5b_LLM_vs_Det)."""
    path = os.path.join(out_dir, "agentic_eval_results.xlsx")
    if not os.path.exists(path):
        return None
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    if "T5b_LLM_vs_Det" not in wb.sheetnames:
        return None
    rows = list(wb["T5b_LLM_vs_Det"].iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if r and "llm_wall_s_per_session" in r:
            j = r.index("llm_wall_s_per_session")
            for rr in rows[i + 1:]:
                if rr and rr[0] and str(rr[0]).startswith("LLM") and isinstance(rr[j], (int, float)):
                    return float(rr[j])
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default="test.csv")
    ap.add_argument("--pcap", default=None, help="one pcap (>= 20 s) for the feature-extraction time")
    ap.add_argument("--out-dir", default="reviewer_package")
    ap.add_argument("--n", type=int, default=4000, help="windows used for the inference timing")
    ap.add_argument("--reps", type=int, default=5)
    ap.add_argument("--skip-model", action="store_true", help="skip PatchTST timing (no torch needed)")
    args = ap.parse_args()

    app = load_module("agentic_app5", os.path.join(HERE, "agentic_app5.py"))
    res = {}
    vals = {}
    if not args.skip_model:
        pred = load_module("user_predict", app.PREDICT_SCRIPT)
        print("timing PatchTST inference …")
        res["inference"] = measure_inference(pred, args.csv, args.n, args.reps)
        i = res["inference"]
        vals["latInf"] = f"{i['ms_per_window_batched']:.2f}\\,ms ({i['windows_per_s']:.0f}\\,win/s)"
        if args.pcap:
            print("timing feature extraction …")
            res["features"] = measure_features(pred, args.pcap)
            vals["latFeat"] = f"{res['features']['s_per_20s_recording']:.2f}\\,s"
    c = measure_controller(app, args.out_dir)
    if c:
        res["controller"] = c
        vals["latCtrlE"] = f"{c['evidence_ms']:.1f}\\,ms"
        vals["latCtrlF"] = f"{c['fast_ms']:.1f}\\,ms"
    llm = read_llm_seconds(args.out_dir)
    if llm is not None:
        res["llm_s_per_recording"] = llm
        vals["latLLM"] = f"{llm:.0f}\\,s"
    with open("latency_results.json", "w", encoding="utf-8") as f:
        json.dump(res, f, indent=1)
    with open("latency_values.tex", "w", encoding="utf-8") as f:
        f.write("% generated by measure_latency.py — copy next to the paper .tex\n")
        for k, v in vals.items():
            f.write("\\renewcommand{\\%s}{%s}\n" % (k, v))
    print(json.dumps(res, indent=1))
    print("\nWrote latency_values.tex and latency_results.json.")
    print("Copy latency_values.tex into the folder of the paper .tex and recompile.")
    if "features" not in res:
        print("(no --pcap given: feature-extraction cell stays 'n/m')")


if __name__ == "__main__":
    main()
